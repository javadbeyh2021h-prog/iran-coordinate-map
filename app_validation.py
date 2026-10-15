from flask import Flask, render_template_string, request, send_file, jsonify
from io import BytesIO
import random
import time
import re
import requests
from openpyxl import load_workbook, Workbook

app = Flask(__name__)

HTML = r'''<!DOCTYPE html>
<html lang="fa" dir="rtl">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>راه‌بر | نقشه مسیر و باربرگ</title>
  <link rel="stylesheet" href="https://unpkg.com/leaflet@1.9.4/dist/leaflet.css">
  <style>
    :root{--bg:#f3f6fb;--panel:#fff;--text:#172033;--muted:#64748b;--line:#dce4ef;--primary:#2457d6;--green:#13834b;--danger:#c52b3a;--soft:#edf3ff;--shadow:0 12px 35px #20365b12}
    body.dark{--bg:#101827;--panel:#182235;--text:#edf3ff;--muted:#a7b4ca;--line:#34435a;--primary:#7aa2ff;--green:#40cb87;--danger:#ff7581;--soft:#243450;--shadow:0 12px 35px #0003}
    *{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--text);font-family:Tahoma,Arial,sans-serif;transition:background .2s,color .2s}header{background:linear-gradient(120deg,#143b9b,#2869e8);color:white;padding:22px 16px}header .head{max-width:1200px;margin:auto;display:flex;align-items:center;justify-content:space-between;gap:12px;flex-wrap:wrap}h1{font-size:24px;margin:0 0 7px}header p{margin:0;color:#e3ebff;font-size:13px;line-height:1.8}.head-actions{display:flex;gap:8px;flex-wrap:wrap}.head-actions button{background:#ffffff20;border:1px solid #ffffff55}.wrap{max-width:1200px;width:calc(100% - 24px);margin:18px auto 35px}.panel{background:var(--panel);border:1px solid var(--line);border-radius:16px;padding:16px;margin-bottom:15px;box-shadow:var(--shadow)}.section-title{font-size:16px;font-weight:bold;margin:0 0 13px}.grid2{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:12px}.grid3{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:12px}.field-card,.metric,.subpanel{background:var(--bg);border:1px solid var(--line);padding:14px;border-radius:12px;min-width:0}.field-card strong{display:block;margin-bottom:7px}.field-card input,.field-card select,.field-card textarea,.form-grid input,.form-grid select{width:100%;padding:11px 12px;border:1px solid var(--line);border-radius:9px;background:var(--panel);color:var(--text);font:inherit;font-size:13px;margin-top:7px}.buttons{display:flex;flex-wrap:wrap;gap:7px;margin-top:9px}button,.file-label{border:0;border-radius:9px;padding:10px 13px;background:var(--primary);color:white;font:inherit;font-size:13px;cursor:pointer;transition:filter .15s}button:hover,.file-label:hover{filter:brightness(.92)}button.green{background:var(--green)}button.gray{background:#596579}button.red{background:var(--danger)}button.outline{background:transparent;color:var(--text);border:1px solid var(--line)}button:disabled{opacity:.5;cursor:not-allowed}#map{height:54vh;min-height:360px;border-radius:13px;border:1px solid var(--line);overflow:hidden}.status{line-height:1.9;color:var(--muted);font-size:13px;margin:10px 0}.coords{direction:ltr;text-align:left;font-weight:bold;overflow-wrap:anywhere;margin:6px 0 10px;font-size:14px}.metric strong{display:block;font-size:13px;margin-bottom:9px}.big{font-size:22px;font-weight:bold;color:var(--primary)}.small{font-size:12px;line-height:1.8;color:var(--muted)}.tag{display:inline-block;padding:4px 8px;border-radius:30px;background:var(--soft);color:var(--primary);font-size:11px;margin-bottom:8px}.route-list{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:10px}.saved-item{padding:12px;border:1px solid var(--line);border-radius:11px;background:var(--bg);min-width:0}.saved-item .name{font-weight:bold;margin-bottom:7px}.saved-item .small{margin:4px 0}.saved-item .buttons button{padding:8px 10px}.form-grid{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:10px}.empty{padding:18px;text-align:center;color:var(--muted);border:1px dashed var(--line);border-radius:10px;font-size:13px}.footer{text-align:center;color:var(--muted);font-size:11px;padding:12px}.hidden{display:none!important}
    @media(max-width:720px){h1{font-size:20px}.grid2,.grid3,.route-list,.form-grid{grid-template-columns:1fr}#map{height:50vh;min-height:320px}.panel{padding:12px}.wrap{width:calc(100% - 16px);margin-top:10px}.big{font-size:20px}}

    .iran-banner{display:block;width:100%;max-width:1200px;max-height:340px;object-fit:cover;object-position:center;margin:20px auto;border-radius:18px;box-shadow:0 8px 25px rgba(0,0,0,.2)}
    @media(max-width:650px){.iran-banner{max-height:220px;margin:12px auto;border-radius:12px}}
  </style>
</head>
<body>
<header><div class="head"><div><h1>🧭 راه‌بر | نقشه مسیر ایران</h1><p>انتخاب مبدأ و مقصد، محاسبه مسیر جاده‌ای، ذخیره مسیرها و آماده‌سازی اطلاعات باربرگ</p></div><div class="head-actions"><button onclick="toggleTheme()" id="themeBtn">🌙 حالت شب</button><button onclick="exportCSV()">⬇ خروجی اکسل (CSV)</button></div></div></header>
<div class="wrap"><img src="/iran-map.png" alt="نقشه سه‌بعدی ایران" class="iran-banner"></div>
<main class="wrap">
  <section class="panel">
    <div class="section-title">۱. انتخاب مبدأ و مقصد</div>
    <div class="grid2">
      <div class="field-card"><strong>📍 مبدأ</strong><input id="originSearch" placeholder="مثلاً لای‌بید، اصفهان" autocomplete="off"><div class="buttons"><button onclick="searchPlace('origin')">جست‌وجوی مبدأ</button><button class="gray" onclick="selectMode('origin')">انتخاب روی نقشه</button></div><div class="small">جست‌وجو با نام شهر، روستا یا مکان</div></div>
      <div class="field-card"><strong>🏁 مقصد</strong><input id="destinationSearch" placeholder="مثلاً اصفهان" autocomplete="off"><div class="buttons"><button class="green" onclick="searchPlace('destination')">جست‌وجوی مقصد</button><button class="gray" onclick="selectMode('destination')">انتخاب روی نقشه</button></div><div class="small">برای انتخاب دستی، دکمه را بزن و روی نقشه کلیک کن</div></div>
    </div>
    <p class="status" id="status">برای شروع، نام مکان را جست‌وجو کن یا حالت انتخاب روی نقشه را فعال کن.</p>
    <div id="map"></div><p class="small">نقشه و جست‌وجو به اینترنت نیاز دارند. مختصات و مسیرهای ذخیره‌شده در همین مرورگر نگهداری می‌شوند.</p>
  </section>
  <section class="panel"><div class="section-title">۲. مختصات و اطلاعات مسیر</div>
    <div class="grid2">
      <div class="field-card"><span class="tag">مبدأ</span><div>Latitude / عرض جغرافیایی</div><div class="coords" id="originLat">—</div><div>Longitude / طول جغرافیایی</div><div class="coords" id="originLng">—</div><button onclick="copyPoint('origin')">کپی مختصات مبدأ</button></div>
      <div class="field-card"><span class="tag">مقصد</span><div>Latitude / عرض جغرافیایی</div><div class="coords" id="destinationLat">—</div><div>Longitude / طول جغرافیایی</div><div class="coords" id="destinationLng">—</div><button class="green" onclick="copyPoint('destination')">کپی مختصات مقصد</button></div>
    </div>
    <div class="grid3" style="margin-top:12px"><div class="metric"><strong>🚗 فاصله جاده‌ای</strong><div class="big" id="routeDistance">—</div></div><div class="metric"><strong>⏱ زمان رانندگی تخمینی</strong><div class="big" id="routeDuration">—</div></div><div class="metric"><strong>🧭 وضعیت</strong><div id="routeStatus" class="small">هنوز مسیری محاسبه نشده است.</div></div></div>
    <div class="buttons"><button onclick="calculateRoute()">محاسبه مسیر جاده‌ای</button><button class="green" onclick="openGoogleMaps()">بازکردن در Google Maps</button><button class="gray" onclick="copyBoth()">کپی هر دو مختصات</button><button class="red" onclick="clearAll()">پاک‌کردن مبدأ و مقصد</button></div>
  </section>
  <section class="panel" id="validationPanel"><div class="section-title">۳. بررسی اولیه باربرگ شهری</div><p class="small">بررسی اولیه نام محدوده از نقشه عمومی و فاصله مسیر است؛ مرز رسمی شهرداری و تأیید سامانه باربرگ را تضمین نمی‌کند. اگر اطلاعات کافی نباشد، نتیجه «نیازمند بررسی» خواهد بود.</p><div class="buttons"><button onclick="validateCurrentRoute()">بررسی مبدأ و مقصد فعلی</button></div><p class="status" id="validationStatus">ابتدا مبدأ و مقصد را انتخاب کن.</p></section>
  <section class="panel"><div class="section-title">۳. ذخیره مسیرهای پرتکرار</div><div class="grid2"><div class="field-card"><strong>نام دلخواه مسیر</strong><input id="routeName" placeholder="مثلاً لای‌بید به اصفهان"><div class="buttons"><button onclick="saveRoute()">ذخیره مسیر فعلی</button></div><div class="small">ابتدا مبدأ و مقصد را انتخاب کن؛ بعد مسیر را ذخیره کن.</div></div><div class="field-card"><strong>مدیریت مسیرها</strong><div class="buttons"><button class="outline" onclick="exportCSV()">خروجی CSV از مسیرهای ذخیره‌شده</button><button class="red" onclick="clearSavedRoutes()">حذف همه مسیرهای ذخیره‌شده</button></div><div class="small">مسیرها فقط در مرورگر همین دستگاه ذخیره می‌شوند؛ با پاک‌کردن داده‌های مرورگر ممکن است حذف شوند.</div></div></div><div id="savedRoutes" class="route-list" style="margin-top:12px"></div></section>

  <section class="panel">
    <div class="section-title">۵. بررسی مختصات واقعی باربرگ در اکسل</div>
    <p class="small">فایل باربرگ را انتخاب کن. مختصات موجود در ستون‌های I تا L تغییر نمی‌کنند. برنامه فاصله مستقیم را محاسبه و ستون‌های نتیجه و توضیح بررسی را به انتهای فایل اضافه می‌کند. این بررسی جایگزین مرز رسمی شهر یا تأیید سامانه باربرگ نیست.</p>
    <div class="field-card">
      <strong>فایل Excel باربرگ</strong>
      <input id="excelFile" type="file" accept=".xlsx,.xlsm">
      <div class="buttons"><button onclick="processExcel()">بررسی مختصات و دریافت گزارش</button></div>
      <p class="status" id="excelStatus">ستون‌های I تا L: مختصات واقعی مبدأ و مقصد؛ مختصات اصلی حفظ می‌شوند.</p>
      <a id="excelDownload" class="hidden" style="display:inline-block;margin-top:8px;color:var(--primary);font-weight:bold" href="#">دانلود فایل آماده‌شده</a>
    </div>
  </section>
  <section class="panel"><div class="section-title">۴. اطلاعات کمکی باربرگ شهری</div><p class="small">این بخش برای مرتب‌کردن اطلاعات مسیر و خروجی گرفتن است؛ اطلاعات به سامانه رسمی باربرگ ارسال نمی‌شود.</p><div class="form-grid"><label class="small">نام راننده<input id="driverName" placeholder="نام راننده"></label><label class="small">شماره پلاک<input id="plate" placeholder="مثلاً ۱۲الف۳۴۵ ایران ۶۷"></label><label class="small">نوع بار<input id="cargo" placeholder="نوع بار"></label><label class="small">وزن بار (کیلوگرم)<input id="weight" type="number" min="0" placeholder="مثلاً ۱۲۰۰"></label><label class="small">شماره باربرگ (اختیاری)<input id="barbargNo" placeholder="شماره باربرگ"></label><label class="small">توضیحات<input id="notes" placeholder="توضیحات مسیر یا بار"></label></div><div class="buttons"><button onclick="copyBarbargSummary()">کپی خلاصه اطلاعات</button><button class="green" onclick="exportCSV()">خروجی اکسل (CSV)</button></div><p id="barbargStatus" class="status"></p></section>
  <div class="footer">نقشه با Leaflet و داده‌های OpenStreetMap نمایش داده می‌شود. فاصله و زمان جاده‌ای با سرویس عمومی OSRM محاسبه می‌شود و ممکن است موقتاً در دسترس نباشد.</div>
</main>
<script src="https://unpkg.com/leaflet@1.9.4/dist/leaflet.js"></script>
<script>
let map, routeLayer=null, mode=null, routeRequestId=0;
const points={origin:{lat:null,lng:null,marker:null},destination:{lat:null,lng:null,marker:null}};
const STORAGE_KEY='iranCoordinateMapPoints_v2', ROUTES_KEY='iranCoordinateMapSavedRoutes_v1', THEME_KEY='iranCoordinateMapTheme';
function setStatus(text){document.getElementById('status').textContent=text}
function validPoint(p){return Number.isFinite(p.lat)&&Number.isFinite(p.lng)}
function initializeMap(){
  map=L.map('map').setView([32.4,53.7],5);
  L.tileLayer('https://{s}.tile.openstreetmap.org/{z}/{x}/{y}.png',{maxZoom:19,attribution:'&copy; OpenStreetMap contributors'}).addTo(map);
  map.on('click',e=>{if(!mode){setStatus('برای انتخاب نقطه، ابتدا دکمه انتخاب مبدأ یا مقصد روی نقشه را بزن.');return;}setPoint(mode,e.latlng.lat,e.latlng.lng,true);mode=null;});
  if(localStorage.getItem(THEME_KEY)==='dark')document.body.classList.add('dark'); updateThemeButton();
  restorePoints();renderSavedRoutes();
}
function selectMode(which){mode=which;setStatus(which==='origin'?'حالت انتخاب مبدأ فعال شد؛ روی نقشه کلیک کن.':'حالت انتخاب مقصد فعال شد؛ روی نقشه کلیک کن.');}
function setPoint(which,lat,lng,zoom){const p=points[which];if(p.marker)map.removeLayer(p.marker);p.lat=lat;p.lng=lng;p.marker=L.marker([lat,lng]).addTo(map).bindPopup(which==='origin'?'مبدأ':'مقصد');document.getElementById(which+'Lat').textContent=lat.toFixed(6);document.getElementById(which+'Lng').textContent=lng.toFixed(6);if(zoom)map.setView([lat,lng],Math.max(map.getZoom(),10));savePoints();clearRouteDisplay();if(validPoint(points.origin)&&validPoint(points.destination))calculateRoute();}
function savePoints(){try{localStorage.setItem(STORAGE_KEY,JSON.stringify({origin:{lat:points.origin.lat,lng:points.origin.lng},destination:{lat:points.destination.lat,lng:points.destination.lng}}))}catch(e){}}
function restorePoints(){try{const data=JSON.parse(localStorage.getItem(STORAGE_KEY)||'null');if(!data)return;['origin','destination'].forEach(k=>{if(data[k]&&Number.isFinite(data[k].lat)&&Number.isFinite(data[k].lng))setPoint(k,data[k].lat,data[k].lng,false)});if(validPoint(points.origin)&&validPoint(points.destination)){const group=L.featureGroup([points.origin.marker,points.destination.marker]);map.fitBounds(group.getBounds().pad(.2));calculateRoute();}else if(validPoint(points.origin))map.setView([points.origin.lat,points.origin.lng],10);else if(validPoint(points.destination))map.setView([points.destination.lat,points.destination.lng],10);}catch(e){}}
async function searchPlace(which){const input=document.getElementById(which==='origin'?'originSearch':'destinationSearch');const query=input.value.trim();if(!query){setStatus('نام شهر، روستا یا مکان را وارد کن.');input.focus();return;}setStatus('در حال جست‌وجوی مکان...');try{const url='https://nominatim.openstreetmap.org/search?format=jsonv2&limit=5&countrycodes=ir&q='+encodeURIComponent(query);const response=await fetch(url,{headers:{'Accept-Language':'fa'}});if(!response.ok)throw Error('search failed');const results=await response.json();if(!results.length){setStatus('مکانی پیدا نشد؛ نام دیگری امتحان کن یا روی نقشه انتخاب کن.');return;}let result=results[0];if(results.length>1){const choices=results.map((r,i)=>(i+1)+'. '+r.display_name).join('\n');const answer=prompt('شماره مکان موردنظر را انتخاب کن:\n'+choices,'1');if(answer===null)return;const index=Number(answer)-1;if(!Number.isInteger(index)||index<0||index>=results.length){alert('شماره انتخابی معتبر نیست.');return;}result=results[index];}input.value=result.display_name;setPoint(which,Number(result.lat),Number(result.lon),true);setStatus((which==='origin'?'مبدأ':'مقصد')+' پیدا و روی نقشه ثبت شد.');}catch(e){setStatus('جست‌وجو ناموفق بود. اتصال اینترنت یا سرویس جست‌وجو را بررسی کن.');}}
function clearRouteDisplay(){routeRequestId++;if(routeLayer&&map){map.removeLayer(routeLayer);routeLayer=null;}document.getElementById('routeDistance').textContent='—';document.getElementById('routeDuration').textContent='—';document.getElementById('routeStatus').textContent='هنوز مسیری محاسبه نشده است.';}
async function calculateRoute(){if(!validPoint(points.origin)||!validPoint(points.destination)){setStatus('ابتدا مبدأ و مقصد را انتخاب کن.');return;}const requestId=++routeRequestId;document.getElementById('routeStatus').textContent='در حال محاسبه مسیر...';document.getElementById('routeDistance').textContent='...';document.getElementById('routeDuration').textContent='...';try{const start=points.origin.lng+','+points.origin.lat,end=points.destination.lng+','+points.destination.lat;const url='https://router.project-osrm.org/route/v1/driving/'+start+';'+end+'?overview=full&geometries=geojson&steps=false';const response=await fetch(url);if(!response.ok)throw Error('route failed');const data=await response.json();if(requestId!==routeRequestId)return;if(data.code!=='Ok'||!data.routes||!data.routes.length)throw Error('no route');const r=data.routes[0],km=r.distance/1000,mins=Math.round(r.duration/60);document.getElementById('routeDistance').textContent=km.toLocaleString('fa-IR',{maximumFractionDigits:1})+' کیلومتر';document.getElementById('routeDuration').textContent=mins<60?mins.toLocaleString('fa-IR')+' دقیقه':Math.floor(mins/60).toLocaleString('fa-IR')+' ساعت و '+(mins%60).toLocaleString('fa-IR')+' دقیقه';document.getElementById('routeStatus').textContent='مسیر جاده‌ای نمایش داده شد.';routeLayer=L.geoJSON(r.geometry,{style:{color:'#e43e61',weight:5,opacity:.88}}).addTo(map);map.fitBounds(routeLayer.getBounds(),{padding:[25,25]});setStatus('مسیر جاده‌ای محاسبه شد.');}catch(e){if(requestId!==routeRequestId)return;document.getElementById('routeDistance').textContent='در دسترس نیست';document.getElementById('routeDuration').textContent='—';document.getElementById('routeStatus').textContent='سرویس مسیریابی پاسخ نداد؛ دوباره تلاش کن یا Google Maps را باز کن.';setStatus('محاسبه مسیر ناموفق بود؛ ممکن است سرویس عمومی موقتاً در دسترس نباشد.');}}
function openGoogleMaps(){if(!validPoint(points.origin)||!validPoint(points.destination)){alert('ابتدا مبدأ و مقصد را انتخاب کن.');return;}window.open('https://www.google.com/maps/dir/?api=1&origin='+points.origin.lat+','+points.origin.lng+'&destination='+points.destination.lat+','+points.destination.lng+'&travelmode=driving','_blank','noopener');}
async function validateCurrentRoute(){
  const status=document.getElementById('validationStatus');
  if(!validPoint(points.origin)||!validPoint(points.destination)){status.textContent='ابتدا مبدأ و مقصد را انتخاب کن.';return;}
  status.textContent='در حال بررسی نام محدوده و فاصله مسیر...';
  try{
    const response=await fetch('/validate_route',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({origin:points.origin,destination:points.destination})});
    const data=await response.json();
    if(!response.ok)throw Error(data.error||'بررسی ناموفق بود.');
    status.textContent='نتیجه: '+data.result+' — '+data.reason+(data.city_origin?' نام محدوده مبدأ: '+data.city_origin+'.':'')+(data.city_destination?' نام محدوده مقصد: '+data.city_destination+'.':'')+(data.road_km!==null&&data.road_km!==undefined?' فاصله جاده‌ای تقریبی: '+Number(data.road_km).toLocaleString('fa-IR',{maximumFractionDigits:1})+' کیلومتر.':'')+' این نتیجه تأیید رسمی نیست.';
  }catch(e){status.textContent='نتیجه: نیازمند بررسی. '+e.message+' اتصال اینترنت یا سرویس نقشه را بررسی کن.';}
}
async function copyText(text){try{if(navigator.clipboard&&window.isSecureContext){await navigator.clipboard.writeText(text);alert('کپی شد.');}else throw Error('clipboard unavailable');}catch(e){prompt('متن زیر را انتخاب و کپی کن:',text)}}
function copyPoint(which){const p=points[which];if(!validPoint(p)){alert('ابتدا '+(which==='origin'?'مبدأ':'مقصد')+' را انتخاب کن.');return;}copyText(p.lat.toFixed(6)+', '+p.lng.toFixed(6));}
function copyBoth(){if(!validPoint(points.origin)||!validPoint(points.destination)){alert('ابتدا هر دو نقطه را انتخاب کن.');return;}copyText('مبدأ: '+points.origin.lat.toFixed(6)+', '+points.origin.lng.toFixed(6)+'\nمقصد: '+points.destination.lat.toFixed(6)+', '+points.destination.lng.toFixed(6));}
function clearAll(){if(!confirm('مبدأ و مقصد فعلی پاک شوند؟'))return;['origin','destination'].forEach(k=>{const p=points[k];if(p.marker)map.removeLayer(p.marker);p.marker=null;p.lat=null;p.lng=null;document.getElementById(k+'Lat').textContent='—';document.getElementById(k+'Lng').textContent='—';document.getElementById(k==='origin'?'originSearch':'destinationSearch').value='';});localStorage.removeItem(STORAGE_KEY);clearRouteDisplay();setStatus('مبدأ و مقصد پاک شدند.');}
function getSavedRoutes(){try{return JSON.parse(localStorage.getItem(ROUTES_KEY)||'[]')}catch(e){return []}}
function saveRoute(){if(!validPoint(points.origin)||!validPoint(points.destination)){alert('ابتدا مبدأ و مقصد را انتخاب کن.');return;}const name=document.getElementById('routeName').value.trim()||('مسیر '+(getSavedRoutes().length+1));const routes=getSavedRoutes();routes.unshift({id:Date.now(),name,origin:{lat:points.origin.lat,lng:points.origin.lng,label:document.getElementById('originSearch').value.trim()},destination:{lat:points.destination.lat,lng:points.destination.lng,label:document.getElementById('destinationSearch').value.trim()},created:new Date().toLocaleDateString('fa-IR')});localStorage.setItem(ROUTES_KEY,JSON.stringify(routes));document.getElementById('routeName').value='';renderSavedRoutes();setStatus('مسیر «'+name+'» ذخیره شد.');}
function renderSavedRoutes(){const box=document.getElementById('savedRoutes');if(!box)return;const routes=getSavedRoutes();if(!routes.length){box.innerHTML='<div class="empty" style="grid-column:1/-1">هنوز مسیری ذخیره نشده است.</div>';return;}box.innerHTML='';routes.forEach(r=>{const card=document.createElement('div');card.className='saved-item';const title=document.createElement('div');title.className='name';title.textContent=r.name;const desc=document.createElement('div');desc.className='small';desc.textContent=(r.origin.label||r.origin.lat.toFixed(4)+', '+r.origin.lng.toFixed(4))+' ← '+(r.destination.label||r.destination.lat.toFixed(4)+', '+r.destination.lng.toFixed(4));const date=document.createElement('div');date.className='small';date.textContent='تاریخ ذخیره: '+(r.created||'—');const actions=document.createElement('div');actions.className='buttons';const load=document.createElement('button');load.textContent='بارگذاری مسیر';load.onclick=()=>{document.getElementById('originSearch').value=r.origin.label||'';document.getElementById('destinationSearch').value=r.destination.label||'';setPoint('origin',r.origin.lat,r.origin.lng,false);setPoint('destination',r.destination.lat,r.destination.lng,false);map.fitBounds(L.featureGroup([points.origin.marker,points.destination.marker]).getBounds().pad(.2));calculateRoute();setStatus('مسیر ذخیره‌شده بارگذاری شد.');};const del=document.createElement('button');del.className='red';del.textContent='حذف';del.onclick=()=>{localStorage.setItem(ROUTES_KEY,JSON.stringify(getSavedRoutes().filter(x=>x.id!==r.id)));renderSavedRoutes();};actions.append(load,del);card.append(title,desc,date,actions);box.append(card);});}
function clearSavedRoutes(){if(!confirm('تمام مسیرهای ذخیره‌شده حذف شوند؟'))return;localStorage.removeItem(ROUTES_KEY);renderSavedRoutes();}
function csvEscape(value){const s=String(value==null?'':value);return '"'+s.replace(/"/g,'""')+'"';}
function exportCSV(){const routes=getSavedRoutes();const rows=[['نام مسیر','مبدأ','مبدأ Latitude','مبدأ Longitude','مقصد','مقصد Latitude','مقصد Longitude','فاصله جاده‌ای','نام راننده','پلاک','نوع بار','وزن کیلوگرم','شماره باربرگ','توضیحات','تاریخ ذخیره']];const dist=document.getElementById('routeDistance').textContent;const extra={driver:document.getElementById('driverName').value,plate:document.getElementById('plate').value,cargo:document.getElementById('cargo').value,weight:document.getElementById('weight').value,no:document.getElementById('barbargNo').value,notes:document.getElementById('notes').value};if(routes.length){routes.forEach(r=>rows.push([r.name,r.origin.label,r.origin.lat,r.origin.lng,r.destination.label,r.destination.lat,r.destination.lng,'',extra.driver,extra.plate,extra.cargo,extra.weight,extra.no,extra.notes,r.created]));}else if(validPoint(points.origin)&&validPoint(points.destination)){rows.push([document.getElementById('routeName').value||'مسیر فعلی',document.getElementById('originSearch').value,points.origin.lat,points.origin.lng,document.getElementById('destinationSearch').value,points.destination.lat,points.destination.lng,dist,extra.driver,extra.plate,extra.cargo,extra.weight,extra.no,extra.notes,new Date().toLocaleDateString('fa-IR')]);}else{alert('برای خروجی، ابتدا یک مسیر انتخاب کن یا مسیرهای ذخیره‌شده داشته باش.');return;}const csv='\uFEFF'+rows.map(row=>row.map(csvEscape).join(',')).join('\r\n');const blob=new Blob([csv],{type:'text/csv;charset=utf-8;'});const url=URL.createObjectURL(blob);const a=document.createElement('a');a.href=url;a.download='iran-coordinate-routes.csv';document.body.appendChild(a);a.click();a.remove();URL.revokeObjectURL(url);}
function copyBarbargSummary(){if(!validPoint(points.origin)||!validPoint(points.destination)){alert('ابتدا مبدأ و مقصد را انتخاب کن.');return;}const lines=['خلاصه اطلاعات مسیر','مبدأ: '+document.getElementById('originSearch').value,'مبدأ مختصات: '+points.origin.lat.toFixed(6)+', '+points.origin.lng.toFixed(6),'مقصد: '+document.getElementById('destinationSearch').value,'مقصد مختصات: '+points.destination.lat.toFixed(6)+', '+points.destination.lng.toFixed(6),'فاصله جاده‌ای: '+document.getElementById('routeDistance').textContent,'نام راننده: '+document.getElementById('driverName').value,'پلاک: '+document.getElementById('plate').value,'نوع بار: '+document.getElementById('cargo').value,'وزن کیلوگرم: '+document.getElementById('weight').value,'شماره باربرگ: '+document.getElementById('barbargNo').value,'توضیحات: '+document.getElementById('notes').value];copyText(lines.join('\n'));document.getElementById('barbargStatus').textContent='خلاصه اطلاعات آماده کپی است. این اطلاعات به سامانه رسمی ارسال نشده است.';}
function toggleTheme(){document.body.classList.toggle('dark');localStorage.setItem(THEME_KEY,document.body.classList.contains('dark')?'dark':'light');updateThemeButton()}
function updateThemeButton(){document.getElementById('themeBtn').textContent=document.body.classList.contains('dark')?'☀️ حالت روز':'🌙 حالت شب'}

async function processExcel(){
  const file=document.getElementById('excelFile').files[0];
  const status=document.getElementById('excelStatus'), link=document.getElementById('excelDownload');
  link.classList.add('hidden');
  if(!file){status.textContent='اول فایل اکسل را انتخاب کن.';return;}
  const form=new FormData();form.append('file',file);
  status.textContent='در حال پردازش فایل؛ ممکن است کمی طول بکشد...';
  try{
    const response=await fetch('/process_excel',{method:'POST',body:form});
    if(!response.ok){let msg='پردازش ناموفق بود.';try{const j=await response.json();msg=j.error||msg}catch(e){}throw Error(msg)}
    const blob=await response.blob(), url=URL.createObjectURL(blob);
    link.href=url;link.download='barbarg_coordinate_review.xlsx';link.classList.remove('hidden');
    status.textContent='گزارش آماده شد. مختصات و ستون‌های اصلی حفظ شده‌اند و نتیجه بررسی به انتهای فایل اضافه شده است.';
  }catch(e){status.textContent='خطا: '+e.message;}
}

window.addEventListener('load',initializeMap);
</script>
</body></html>'''

@app.route('/')
def home():
    return render_template_string(HTML)


@app.route('/iran-map.png')
def iran_map_image():
    return send_file(
        app.root_path + '/iran-map.png',
        mimetype='image/png'
    )


# Province code mapping based on commonly published Iranian private-car plate code tables.
# Some codes can overlap by letter/region; this tool only sees the two-digit code,
# so ambiguous/changed allocations should be verified against the plate letter when needed.
PLATE_PROVINCES = {
    # Tehran / Alborz
    10:"تهران",11:"تهران",20:"تهران",22:"تهران",30:"البرز",33:"تهران",38:"البرز",40:"تهران",44:"تهران",50:"تهران",55:"تهران",60:"تهران",66:"تهران",68:"البرز",70:"تهران",77:"تهران",78:"البرز",80:"تهران",88:"تهران",90:"تهران",99:"تهران",
    # Isfahan
    13:"اصفهان",23:"اصفهان",43:"اصفهان",53:"اصفهان",67:"اصفهان",
    # Razavi / North / South Khorasan (some codes require plate letter to distinguish)
    12:"خراسان رضوی",32:"خراسان رضوی",36:"خراسان رضوی",42:"خراسان رضوی",74:"خراسان رضوی",26:"خراسان شمالی",52:"خراسان جنوبی",
    # Khuzestan and East/West Azerbaijan
    14:"خوزستان",24:"خوزستان",34:"خوزستان",15:"آذربایجان شرقی",25:"آذربایجان شرقی",35:"آذربایجان شرقی",17:"آذربایجان غربی",27:"آذربایجان غربی",37:"آذربایجان غربی",
    # Fars, Mazandaran, Kerman, Gilan
    63:"فارس",73:"فارس",83:"فارس",93:"فارس",62:"مازندران",72:"مازندران",82:"مازندران",92:"مازندران",45:"کرمان",65:"کرمان",75:"کرمان",46:"گیلان",56:"گیلان",76:"گیلان",
    # Other provinces
    18:"همدان",28:"همدان",19:"کرمانشاه",29:"کرمانشاه",39:"کرمانشاه",31:"لرستان",41:"لرستان",51:"کردستان",61:"کردستان",47:"مرکزی",57:"مرکزی",48:"بوشهر",58:"بوشهر",71:"چهارمحال و بختیاری",81:"چهارمحال و بختیاری",79:"قزوین",89:"قزوین",84:"هرمزگان",94:"هرمزگان",85:"سیستان و بلوچستان",95:"سیستان و بلوچستان",86:"سمنان",96:"سمنان",87:"زنجان",97:"زنجان",91:"اردبیل",98:"ایلام",49:"کهگیلویه و بویراحمد",59:"گلستان",69:"گلستان",54:"یزد",64:"یزد",16:"قم"
}
# Approximate city centers and sampling envelopes. These are seed areas, not official city boundaries.
# A wider envelope allows finding two points at least 20 km apart, but points are approximate.
PROVINCE_CITIES = {
    "تهران":[("تهران",35.6892,51.3890,.20),("ورامین",35.3252,51.6457,.18),("شهریار",35.6590,51.0590,.18)],
    "البرز":[("کرج",35.8400,50.9391,.18),("نظرآباد",35.9520,50.6070,.16)],
    "اصفهان":[("اصفهان",32.6546,51.6680,.20),("کاشان",33.9850,51.4100,.18),("نجف‌آباد",32.6324,51.3668,.16)],
    "خراسان رضوی":[("مشهد",36.2972,59.6067,.20),("نیشابور",36.2133,58.7958,.18),("سبزوار",36.2126,57.6819,.18)],
    "خراسان شمالی":[("بجنورد",37.4747,57.3290,.18),("شیروان",37.3967,57.9290,.16)],
    "خراسان جنوبی":[("بیرجند",32.8663,59.2211,.18),("قائن",33.7265,59.1844,.16)],
    "خوزستان":[("اهواز",31.3183,48.6706,.20),("دزفول",32.3831,48.4236,.18),("آبادان",30.3473,48.2934,.16)],
    "آذربایجان شرقی":[("تبریز",38.0962,46.2738,.20),("مراغه",37.3892,46.2374,.16),("میانه",37.4210,47.7150,.16)],
    "آذربایجان غربی":[("ارومیه",37.5527,45.0761,.18),("خوی",38.5503,44.9521,.16),("مهاباد",36.7631,45.7222,.16)],
    "فارس":[("شیراز",29.5918,52.5837,.20),("مرودشت",29.8742,52.8025,.18),("کازرون",29.6195,51.6541,.16)],
    "مازندران":[("ساری",36.5659,53.0586,.18),("آمل",36.4696,52.3507,.16),("بابل",36.5387,52.6765,.16)],
    "کرمان":[("کرمان",30.2839,57.0834,.20),("رفسنجان",30.4067,55.9939,.18),("سیرجان",29.4514,55.6800,.18)],
    "گیلان":[("رشت",37.2808,49.5832,.18),("لاهیجان",37.2072,50.0039,.16),("انزلی",37.4727,49.4622,.16)],
    "همدان":[("همدان",34.7981,48.5150,.18),("ملایر",34.2969,48.8235,.16),("نهاوند",34.1885,48.3768,.16)],
    "کرمانشاه":[("کرمانشاه",34.3142,47.0650,.18),("اسلام‌آباد غرب",34.1094,46.5275,.16),("کنگاور",34.5043,47.9653,.16)],
    "لرستان":[("خرم‌آباد",33.4878,48.3558,.18),("بروجرد",33.8942,48.7517,.16)],
    "کردستان":[("سنندج",35.3111,46.9964,.18),("سقز",36.2499,46.2735,.16)],
    "مرکزی":[("اراک",34.0917,49.6892,.18),("ساوه",35.0213,50.3566,.16)],
    "بوشهر":[("بوشهر",28.9234,50.8203,.18),("برازجان",29.2666,51.2159,.16)],
    "چهارمحال و بختیاری":[("شهرکرد",32.3256,50.8644,.16),("بروجن",31.9652,51.2873,.16)],
    "قزوین":[("قزوین",36.2688,50.0041,.18),("تاکستان",36.0697,49.6959,.16)],
    "هرمزگان":[("بندرعباس",27.1832,56.2666,.18),("میناب",27.1467,57.0750,.16)],
    "سیستان و بلوچستان":[("زاهدان",29.4963,60.8629,.20),("زابل",31.0306,61.4949,.18)],
    "سمنان":[("سمنان",35.5769,53.3950,.18),("شاهرود",36.4182,54.9763,.16)],
    "زنجان":[("زنجان",36.6736,48.4787,.18),("ابهر",36.1468,49.2180,.16)],
    "اردبیل":[("اردبیل",38.2498,48.2933,.18),("پارس‌آباد",39.6482,47.9174,.16)],
    "ایلام":[("ایلام",33.6374,46.4227,.18),("دهلران",32.6941,47.2679,.16)],
    "کهگیلویه و بویراحمد":[("یاسوج",30.6682,51.5879,.16),("دهدشت",30.7942,50.5645,.16)],
    "گلستان":[("گرگان",36.8456,54.4393,.18),("گنبدکاووس",37.2500,55.1672,.16)],
    "یزد":[("یزد",31.8974,54.3569,.20),("میبد",32.2487,54.0165,.16)],
    "قم":[("قم",34.6416,50.8746,.18)]
}
def normalize_digits(value):
    return str(value or '').translate(str.maketrans('۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩','01234567890123456789')).strip()
def straight_distance_km(a, b):
    # Haversine straight-line distance (not road distance).
    from math import radians, sin, cos, asin, sqrt
    lat1, lon1 = map(radians, a)
    lat2, lon2 = map(radians, b)
    dlat, dlon = lat2 - lat1, lon2 - lon1
    h = sin(dlat / 2) ** 2 + cos(lat1) * cos(lat2) * sin(dlon / 2) ** 2
    return 6371.0 * 2 * asin(sqrt(min(1.0, h)))

def make_pair(city, minimum_km=20):
    name, lat, lon, radius = city
    # Approximate random points around a city center; no external routing call.
    for _ in range(300):
        a = (lat + random.uniform(-radius, radius), lon + random.uniform(-radius, radius))
        b = (lat + random.uniform(-radius, radius), lon + random.uniform(-radius, radius))
        distance = straight_distance_km(a, b)
        if distance >= minimum_km:
            return a, b, distance, name
    return None

def coordinate_pair_distance_km(lat1, lon1, lat2, lon2):
    from math import radians, sin, cos, asin, sqrt
    lat1, lon1, lat2, lon2 = map(radians, [lat1, lon1, lat2, lon2])
    h = sin((lat2-lat1)/2)**2 + cos(lat1)*cos(lat2)*sin((lon2-lon1)/2)**2
    return 6371.0 * 2 * asin(sqrt(min(1.0, h)))


def valid_coordinate_pair(lat, lon):
    try:
        lat, lon = float(lat), float(lon)
        return -90 <= lat <= 90 and -180 <= lon <= 180
    except (TypeError, ValueError):
        return False


def roughly_in_iran(lat, lon):
    # Broad geographic screening box only; not a legal or administrative boundary.
    return 24.0 <= float(lat) <= 40.5 and 43.0 <= float(lon) <= 64.5


def reverse_place(lat, lon):
    response = requests.get(
        'https://nominatim.openstreetmap.org/reverse',
        params={'format':'jsonv2', 'lat':lat, 'lon':lon, 'zoom':10, 'addressdetails':1},
        headers={'User-Agent':'RahbarIranRouteChecker/1.0 (preliminary route screening)'},
        timeout=8
    )
    response.raise_for_status()
    data = response.json()
    address = data.get('address') or {}
    city = (address.get('city') or address.get('town') or address.get('municipality') or
            address.get('city_district') or address.get('village') or address.get('county') or '')
    return city


@app.route('/validate_route', methods=['POST'])
def validate_route():
    payload = request.get_json(silent=True) or {}
    try:
        origin = payload.get('origin') or {}
        destination = payload.get('destination') or {}
        olat, olon = float(origin['lat']), float(origin['lng'])
        dlat, dlon = float(destination['lat']), float(destination['lng'])
    except (KeyError, TypeError, ValueError):
        return jsonify(error='مختصات مبدأ و مقصد معتبر نیستند.'), 400
    if not (valid_coordinate_pair(olat, olon) and valid_coordinate_pair(dlat, dlon)):
        return jsonify(error='مقادیر مختصات خارج از بازه مجاز هستند.'), 400
    if not (roughly_in_iran(olat, olon) and roughly_in_iran(dlat, dlon)):
        return jsonify(result='ناسازگار از نظر موقعیت', reason='حداقل یکی از نقاط خارج از محدوده جغرافیایی تقریبی ایران است؛ این محدوده فقط غربالگری کلی است.', city_origin='', city_destination='', road_km=None)
    city_origin = city_destination = ''
    road_km = None
    try:
        city_origin = reverse_place(olat, olon)
        time.sleep(1.1)  # Respect public Nominatim rate limits.
        city_destination = reverse_place(dlat, dlon)
    except requests.RequestException:
        pass
    try:
        route_response = requests.get(
            f'https://router.project-osrm.org/route/v1/driving/{olon},{olat};{dlon},{dlat}',
            params={'overview':'false'}, timeout=8
        )
        route_response.raise_for_status()
        route_data = route_response.json()
        if route_data.get('code') == 'Ok' and route_data.get('routes'):
            road_km = round(route_data['routes'][0]['distance'] / 1000, 2)
    except (requests.RequestException, ValueError, KeyError, TypeError):
        pass
    if city_origin and city_destination and city_origin == city_destination:
        result = 'سازگار اولیه'
        reason = 'سرویس آدرس عمومی هر دو نقطه را به یک نام محدوده نسبت داده است؛ این اثبات حضور داخل مرز رسمی شهر نیست.'
    elif city_origin and city_destination:
        result = 'نیازمند بررسی'
        reason = 'نام محدوده‌های برگشتی متفاوت‌اند؛ مرز رسمی و شرایط باربرگ باید بررسی شود.'
    else:
        result = 'نیازمند بررسی'
        reason = 'نام محدوده یکی از نقاط از سرویس نقشه قابل تشخیص نبود.'
    return jsonify(result=result, reason=reason, city_origin=city_origin, city_destination=city_destination, road_km=road_km)


@app.route('/process_excel', methods=['POST'])
def process_excel():
    if 'file' not in request.files or not request.files['file'].filename:
        return jsonify(error='فایل اکسل انتخاب نشده است.'), 400
    file = request.files['file']
    filename = file.filename.lower()
    if not filename.endswith(('.xlsx', '.xlsm')):
        return jsonify(error='لطفاً فایل xlsx یا xlsm انتخاب کن.'), 400
    try:
        keep_vba = filename.endswith('.xlsm')
        source_wb = load_workbook(file, data_only=False, keep_vba=keep_vba)
        source_ws = source_wb.active
        # Keep all original columns and coordinates intact; add a transparent audit at the end.
        headers = [str(source_ws.cell(1, c).value or '').strip() for c in range(1, source_ws.max_column + 1)]
        status_col = source_ws.max_column + 1
        distance_col = status_col + 1
        reason_col = status_col + 2
        source_ws.cell(1, status_col, 'نتیجه بررسی اولیه مختصات')
        source_ws.cell(1, distance_col, 'فاصله مستقیم (کیلومتر)')
        source_ws.cell(1, reason_col, 'دلیل / توضیح بررسی')
        checked = 0
        for row in range(2, source_ws.max_row + 1):
            values = [source_ws.cell(row, c).value for c in range(9, 13)]  # I, J, K, L
            try:
                olat, olon, dlat, dlon = [float(v) for v in values]
                coords_ok = (valid_coordinate_pair(olat, olon) and valid_coordinate_pair(dlat, dlon))
            except (TypeError, ValueError):
                coords_ok = False
            if not coords_ok:
                source_ws.cell(row, status_col, 'ناسازگار از نظر مختصات')
                source_ws.cell(row, distance_col, None)
                source_ws.cell(row, reason_col, 'یک یا چند مختصات خالی یا خارج از بازه معتبر است؛ مختصات اصلی تغییر داده نشد.')
                continue
            distance = round(coordinate_pair_distance_km(olat, olon, dlat, dlon), 2)
            source_ws.cell(row, distance_col, distance)
            checked += 1
            if not (roughly_in_iran(olat, olon) and roughly_in_iran(dlat, dlon)):
                source_ws.cell(row, status_col, 'نیازمند بررسی')
                source_ws.cell(row, reason_col, 'حداقل یک نقطه خارج از محدوده جغرافیایی تقریبی ایران است؛ این آزمون مرز رسمی نیست. مختصات اصلی حفظ شد.')
            else:
                source_ws.cell(row, status_col, 'نیازمند بررسی')
                note = 'مختصات از نظر عددی قابل‌خواندن و در محدوده جغرافیایی تقریبی ایران هستند؛ برای تأیید داخل محدوده رسمی شهر باید مرز معتبر شهر بررسی شود.'
                if distance < 20:
                    note += ' فاصله مستقیم کمتر از ۲۰ کیلومتر است؛ این مورد به‌تنهایی به معنی نامعتبر بودن باربرگ نیست.'
                note += ' فاصله ثبت‌شده مستقیم است، نه جاده‌ای؛ مختصات اصلی تغییر نکرد.'
                source_ws.cell(row, reason_col, note)
        source_wb.calculation.fullCalcOnLoad = True
        output = BytesIO()
        source_wb.save(output)
        source_wb.close()
        output.seek(0)
        extension = '.xlsm' if keep_vba else '.xlsx'
        outname = 'barbarg_coordinate_review' + extension
        mimetype = ('application/vnd.ms-excel.sheet.macroEnabled.12' if keep_vba else
                    'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        return send_file(output, as_attachment=True, download_name=outname, mimetype=mimetype)
    except Exception as e:
        return jsonify(error='فایل پردازش نشد. مطمئن شو فایل Excel سالم است. جزئیات: ' + str(e)), 500


if __name__ == '__main__':
    app.run(host='127.0.0.1', port=5000, debug=True)
